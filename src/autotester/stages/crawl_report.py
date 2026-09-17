"""The crawl, as something a human can read: an Excel workbook and the
summary rows the crawl page shows (Track B5).

Presentation over evidence already on disk — like `stages/report_export.py`,
this is never a second source of truth. Its job is to make the two things a
bounded crawl must be judged on impossible to miss: **why it stopped**, and
**what it refused to touch**.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from autotester.core.excel import autosize_columns
from autotester.schema.crawl import Crawl, CrawlCoverage, CrawlIssue
from autotester.schema.enums import EdgeOutcome, IssueKind
from autotester.schema.screen_graph import ScreenEdge, ScreenNode
from autotester.stages.explore_status import displayed_status
from autotester.store.project_store import ProjectStore

_REFUSED = (
    EdgeOutcome.DENIED_POLICY,
    EdgeOutcome.SKIPPED_UNNAMED,
    EdgeOutcome.OFF_DOMAIN_REFUSED,
)


def crawl_summary(crawl: Crawl) -> list[tuple[str, str]]:
    """(label, value) rows for the summary sheet and the crawl page. A list of
    pairs rather than a model: this is presentation, not a domain shape (C1)."""
    return [
        ("Crawl", crawl.id),
        ("Project", crawl.project),
        ("Status", displayed_status(crawl).value),  # X18(d): a legacy no-op is not success
        ("Stopped because", crawl.stop_reason or "—"),
        ("Write policy", crawl.policy.write_policy.value),
        ("Screens found", str(crawl.screens)),
        *coverage_rows(crawl),
        ("Actions tried", str(crawl.actions)),
        ("Edges recorded", str(crawl.edges)),
        ("Refused by policy", str(crawl.denied)),
        ("Issues found (in the product)", str(crawl.issues)),
        ("Tool failures (the crawler's own)", str(crawl.tool_failures)),
        ("Started", crawl.started_at or "—"),
        ("Finished", crawl.finished_at or "—"),
    ]


def coverage_figure(crawl: Crawl) -> str:
    """The one-line coverage headline (V7d), or an honest 'not recorded' for an old crawl."""
    cov = crawl.coverage
    if cov is None:
        return "not recorded (this crawl predates coverage)"
    if cov.error:
        return f"could not be computed ({cov.error})"
    return f"{cov.percent}% of controls ({cov.controls_exercised} of {cov.controls_discovered})"


def coverage_rows(crawl: Crawl) -> list[tuple[str, str]]:
    """Coverage beside the stop reason, with the same billing (V7d)."""
    cov = crawl.coverage
    rows = [("Coverage", coverage_figure(crawl))]
    if cov is None:
        return rows
    rows.append(("Screens reached / left queued",
                 f"{cov.screens_reached} / {cov.screens_queued_unvisited}"))
    rows.append(("Not exercised, by reason",
                 "; ".join(f"{r}: {n}" for r, n in cov.by_reason().items()) or "none"))
    rows.append(("Screens not entered, by reason",  # AT-470
                 "; ".join(f"{r}: {n}" for r, n in screens_by_reason(cov).items()) or "none"))
    if cov.spec_error:
        rows.append(("FlowSpec screens reached", cov.spec_error))
    elif cov.spec_screens_total is not None:
        rows.append(("FlowSpec screens reached",
                     f"{cov.spec_screens_reached} of {cov.spec_screens_total}"))
    return rows


def screens_by_reason(cov: CrawlCoverage) -> dict[str, int]:
    counts: dict[str, int] = {}
    for hole in cov.screens_not_entered:
        counts[hole.reason] = counts.get(hole.reason, 0) + 1
    return dict(sorted(counts.items()))


def _screens_sheet(wb: Workbook, nodes: list[ScreenNode]) -> None:
    ws = wb.create_sheet("Screens")
    ws.append(["Screen", "URL template", "Example URL", "Depth", "Status",
               "Controls", "Console errors", "Failed requests"])
    for node in nodes:
        ws.append([
            node.name or node.title or "—", node.url_template, node.url_example,
            node.depth, node.status.value, len(node.elements),
            len(node.console_errors), len(node.failed_requests),
        ])
    autosize_columns(ws)


def _edges_sheet(wb: Workbook, edges: list[ScreenEdge], names: dict[str, str]) -> None:
    ws = wb.create_sheet("Edges")
    ws.append(["From", "Control", "Action", "Outcome", "To", "Reason"])
    for edge in edges:
        ws.append([
            names.get(edge.from_node, edge.from_node), edge.name or edge.target,
            edge.action.value, edge.outcome.value,
            names.get(edge.to_node or "", edge.to_node or "—"), edge.reason or "",
        ])
    autosize_columns(ws)


def _refused_sheet(wb: Workbook, edges: list[ScreenEdge], names: dict[str, str]) -> None:
    """What the crawl deliberately did NOT do. On a real product this is the
    sheet a human reads first: an icon-only button listed here as
    `skipped_unnamed` is a control that needs an `aria-label`."""
    ws = wb.create_sheet("Denied & Skipped")
    ws.append(["Screen", "Control", "Selector", "Refusal", "Reason"])
    for edge in edges:
        if edge.outcome not in _REFUSED:
            continue
        ws.append([
            names.get(edge.from_node, edge.from_node), edge.name or "(unnamed)",
            edge.target, edge.outcome.value, edge.reason or "",
        ])
    autosize_columns(ws)


def _unreached_sheet(wb: Workbook, crawl: Crawl) -> None:
    """Every control the crawl discovered and did not perform, each with its one reason (V7d)."""
    ws = wb.create_sheet("Unreached")
    ws.append(["URL template", "Control", "Selector", "Reason"])
    for hole in (crawl.coverage.holes if crawl.coverage else []):
        ws.append([hole.url_template, hole.name or "(unnamed)", hole.selector, hole.reason])
    for hole in (crawl.coverage.screens_not_entered if crawl.coverage else []):  # AT-470
        ws.append([hole.url_template, f"screen behind '{hole.name or hole.selector}' — not entered",
                   hole.selector, hole.reason])
    autosize_columns(ws)


def _issues_sheet(wb: Workbook, issues: list[CrawlIssue], names: dict[str, str]) -> None:
    ws = wb.create_sheet("Issues")
    ws.append(["Screen", "Kind", "First party", "Detail"])
    for issue in issues:
        ws.append([
            names.get(issue.node_id, issue.node_id), issue.kind.value,
            "yes" if issue.first_party else "no", issue.detail,
        ])
    autosize_columns(ws)


def _tool_failures_sheet(wb: Workbook, issues: list[CrawlIssue],
                         names: dict[str, str]) -> None:
    """What the CRAWLER could not do, kept off the Issues sheet (AT-120).

    Same reasoning as `Noise`: a reader counting rows on "Issues" is counting
    bugs in their product, and a screenshot this tool failed to take is not
    one. Reported rather than dropped so the gap in the evidence is visible."""
    ws = wb.create_sheet("Tool failures")
    ws.append(["Screen", "What the crawler could not do"])
    for issue in issues:
        ws.append([names.get(issue.node_id or "", issue.node_id or "—"), issue.detail])
    autosize_columns(ws)


def _noise_sheet(wb: Workbook, crawl: Crawl) -> None:
    """Third-party hosts whose failures were counted and never reported as
    product issues (X9) — recorded so "we ignored it" is auditable."""
    ws = wb.create_sheet("Noise")
    ws.append(["Third-party host", "Dropped requests"])
    for noise in crawl.noise_counts:
        ws.append([noise.host, noise.count])
    autosize_columns(ws)


def export_crawl_excel(
    project_slug: str, crawl_id: str, out_path: Path, root: Path | None = None
) -> Path:
    """Summary, Screens, Edges, Denied & Skipped, Unreached (V7), Issues, Tool failures, Noise."""
    store = ProjectStore(project_slug, root)
    crawl = store.load_crawl(crawl_id)
    if crawl is None:
        raise ValueError(f"no crawl '{crawl_id}' in project '{project_slug}'")
    nodes = store.list_nodes(crawl_id)
    names = {n.id: (n.name or n.title or n.url_template) for n in nodes}

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Field", "Value"])
    for label, value in crawl_summary(crawl):
        ws.append([label, value])
    autosize_columns(ws)

    edges = store.list_edges(crawl_id)
    _screens_sheet(wb, nodes)
    _edges_sheet(wb, edges, names)
    _refused_sheet(wb, edges, names)
    _unreached_sheet(wb, crawl)
    all_issues = store.list_crawl_issues(crawl_id)
    product = [i for i in all_issues if i.kind is not IssueKind.EVIDENCE]
    tool = [i for i in all_issues if i.kind is IssueKind.EVIDENCE]
    _issues_sheet(wb, product, names)
    _tool_failures_sheet(wb, tool, names)
    _noise_sheet(wb, crawl)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path
