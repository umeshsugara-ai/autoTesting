"""A crawl states its own coverage, and names every hole with one reason (coverage.md V7, AT-459).

Umesh, 2026-09-16: "puura product map hona chiaye na aend to end testing . each possible
route". A report that lists only what a crawl FOUND reads as the whole product; a control never
tried because a bound fired, or never seen because its screen stayed queued, was listed nowhere.
These tests use the scripted fake site (`crawl_fake.py`); the numbers asserted below were read
off a real run of it, not assumed.
"""

from __future__ import annotations

from pathlib import Path

import crawl_fake
import pytest
from crawl_fake import crawl_it
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from autotester.schema.crawl import Crawl, CrawlBounds
from autotester.schema.enums import CrawlStatus, EdgeOutcome, NodeStatus
from autotester.schema.flowspec import FlowSpec, Screen
from autotester.schema.project import Project
from autotester.schema.screen_graph import ElementRef, ScreenEdge, ScreenNode
from autotester.stages.crawl_coverage import REASONS, compute_coverage
from autotester.stages.crawl_report import export_crawl_excel
from autotester.store.project_store import ProjectStore
from autotester.ui.app import app


def _balanced(crawl: Crawl) -> bool:
    """The (c) identity, recounted from the stored holes — never from `percent`."""
    cov = crawl.coverage
    assert cov is not None
    return cov.controls_exercised + sum(cov.by_reason().values()) == cov.controls_discovered


# -- (a)(b)(c): the verify, on a real (fake-site) crawl ----------------------------------

def test_a_crawl_with_a_small_max_actions_states_less_than_full_coverage(tmp_path: Path) -> None:
    """V7's named verify: a deliberately small max_actions -> figure < 100 %, the
    bound:max_actions rows non-empty, and the books balance with at least three reason classes."""
    crawl, _store, _page = crawl_it(tmp_path, bounds=CrawlBounds(max_actions=3))
    cov = crawl.coverage

    assert cov is not None
    assert _balanced(crawl), "(c) the books must balance: exercised + every hole == discovered"
    assert cov.percent < 100
    assert cov.by_reason().get("bound:max_actions", 0) >= 1
    assert len(cov.by_reason()) >= 3, cov.by_reason()
    assert cov.screens_queued_unvisited >= 1 and cov.by_reason().get("not_visited", 0) >= 1


def test_every_hole_has_exactly_one_reason_from_the_closed_set(tmp_path: Path) -> None:
    crawl, _store, _page = crawl_it(tmp_path)
    cov = crawl.coverage

    assert cov is not None and cov.holes
    for hole in cov.holes:
        assert any(hole.reason == r or (r.endswith(":") and hole.reason.startswith(r))
                   for r in REASONS), hole.reason
    assert {"unnamed", "off_domain"} <= set(cov.by_reason())
    assert any(r.startswith("policy:") for r in cov.by_reason())
    assert _balanced(crawl)


def test_the_per_node_cap_is_named_when_it_leaves_controls_untried(tmp_path: Path) -> None:
    crawl, _store, _page = crawl_it(tmp_path, bounds=CrawlBounds(per_node_action_cap=1))

    assert crawl.coverage is not None
    assert crawl.coverage.by_reason().get("bound:per_node_action_cap", 0) >= 1
    assert _balanced(crawl)


def test_a_fully_exercised_completed_crawl_reads_100(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    """The guard must be able to say yes."""
    monkeypatch.setitem(crawl_fake.SITE, crawl_fake.BASE, [
        {"role": "link", "name": "Settings", "selector": "a.settings", "href": "/settings"},
    ])
    monkeypatch.setitem(crawl_fake.SITE, "https://app.test/settings", [])

    crawl, _store, _page = crawl_it(tmp_path)

    assert crawl.status is CrawlStatus.COMPLETED
    assert crawl.coverage is not None and crawl.coverage.percent == 100


# -- unit: the rules a fixture crawl does not easily reach --------------------------------

def _node(status: NodeStatus, *selectors: str) -> ScreenNode:
    return ScreenNode(crawl_id="c", project="demo", url_template="/", url_example="https://x/",
                      signature="s", status=status,
                      elements=[ElementRef(role="button", name=s, selector=s) for s in selectors])


def _compute(status: CrawlStatus, nodes: list[ScreenNode], edges: list[ScreenEdge], *,
             bound: str | None = None, bounds: CrawlBounds | None = None,
             spec: FlowSpec | None = None):
    return compute_coverage(status=status, bound=bound, bounds=bounds or CrawlBounds(),
                            nodes=nodes, edges=edges, spec=spec)


def _edge(node: ScreenNode, target: str, outcome: EdgeOutcome, **extra: object) -> ScreenEdge:
    return ScreenEdge(crawl_id="c", from_node=node.id, target=target, action="click",
                      outcome=outcome, **extra)


def test_a_crawl_that_did_not_complete_never_reads_100() -> None:
    """V7(d): every discovered control performed, but the crawl stopped on a bound."""
    node = _node(NodeStatus.EXPLORED, "#a")

    cov = _compute(CrawlStatus.STOPPED_BOUND, [node], [_edge(node, "#a", EdgeOutcome.SAME_SCREEN)],
                   bound="max_actions")

    assert cov.controls_exercised == cov.controls_discovered == 1
    assert cov.percent < 100


def test_a_refused_submit_on_a_login_wall_is_named_login_wall() -> None:
    node = _node(NodeStatus.EXPLORED, "#go")
    edges = [_edge(node, "#go", EdgeOutcome.DENIED_POLICY, reason="form submit under read_only")]

    assert _compute(CrawlStatus.LOGIN_WALL, [node], edges).by_reason() == {"login_wall": 1}


def test_an_abandoned_screen_names_its_untried_controls_error() -> None:
    cov = _compute(CrawlStatus.COMPLETED, [_node(NodeStatus.ABORTED_ERROR, "#a", "#b")], [])

    assert cov.by_reason() == {"error": 2}


def test_flowspec_screens_reached_are_counted_when_a_spec_exists() -> None:
    spec = FlowSpec(project="demo", screens=[
        Screen(id="s1", name="Home", url_pattern="/"),
        Screen(id="s2", name="Billing", url_pattern="/billing"),
    ])
    cov = _compute(CrawlStatus.COMPLETED, [_node(NodeStatus.EXPLORED)], [], spec=spec)

    assert (cov.spec_screens_reached, cov.spec_screens_total) == (1, 2)


# -- (d): shown on the crawl page and the workbook ------------------------------------------

@pytest.fixture
def stored(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> tuple[Path, Crawl]:
    """A real bounded crawl, then served from the same root the UI reads."""
    monkeypatch.setenv("AUTOTESTER_ROOT", str(tmp_path))
    crawl, store, _page = crawl_it(tmp_path, bounds=CrawlBounds(max_actions=3))
    store.save_project(Project(slug="demo", name="Demo", base_url=crawl_fake.BASE,
                               allowed_domains=["app.test"]))
    return tmp_path, crawl


def test_the_crawl_page_shows_the_coverage_figure_and_reasons(stored: tuple[Path, Crawl]) -> None:
    _root, crawl = stored

    text = TestClient(app).get(f"/projects/demo/crawls/{crawl.id}").text

    tile = (f"<div class='value'>{crawl.coverage.percent}%</div>"
            "<div class='label'>Coverage</div>")
    assert tile in text, "a headline tile beside Screens/Actions, same billing as the stop (V7d)"
    assert f"coverage: {crawl.coverage.percent}% of controls" in text
    assert "bound:max_actions" in text and "not_visited" in text


def test_the_workbook_has_the_figure_and_an_unreached_sheet(
    stored: tuple[Path, Crawl], tmp_path: Path,
) -> None:
    root, crawl = stored

    wb = load_workbook(export_crawl_excel("demo", crawl.id, tmp_path / "c.xlsx", root))
    summary = {r[0]: r[1] for r in wb["Summary"].iter_rows(values_only=True)}
    unreached = list(wb["Unreached"].iter_rows(values_only=True))[1:]

    assert summary["Coverage"].startswith(f"{crawl.coverage.percent}%")
    assert len(unreached) == len(crawl.coverage.holes)
    assert {row[3] for row in unreached} == set(crawl.coverage.by_reason())


def test_a_crawl_recorded_before_coverage_says_so_instead_of_inventing_a_number(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setenv("AUTOTESTER_ROOT", str(tmp_path))
    store = ProjectStore("demo", tmp_path)
    store.save_project(Project(slug="demo", name="Demo", base_url="https://x.test/",
                               allowed_domains=["x.test"]))
    store.save_crawl(Crawl(id="crawl_old", project="demo", status=CrawlStatus.COMPLETED,
                           actions=4, stop_reason="frontier empty"))

    text = TestClient(app).get("/projects/demo/crawls/crawl_old").text

    assert "not recorded" in text
