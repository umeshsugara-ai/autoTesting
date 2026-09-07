"""The crawl Excel report (Track B5). Contract: qa/contracts/explore.md X11 —
the crawl's artifacts must be readable by a human, which includes readable as
the thing testers actually open.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from autotester.schema.crawl import Crawl, CrawlIssue, NoiseCount
from autotester.schema.enums import Action, CrawlStatus, EdgeOutcome, IssueKind
from autotester.schema.screen_graph import ScreenEdge, ScreenNode
from autotester.stages.crawl_report import crawl_summary, export_crawl_excel
from autotester.store.project_store import ProjectStore

SHEETS = ["Summary", "Screens", "Edges", "Denied & Skipped", "Issues", "Noise"]


def seed(tmp_path: Path) -> tuple[ProjectStore, Crawl]:
    store = ProjectStore("erp", tmp_path)
    crawl = Crawl(
        project="erp", id="crawl_demo", status=CrawlStatus.STOPPED_BOUND,
        stop_reason="max_actions", screens=2, actions=9, edges=3, denied=1, issues=1,
        noise_counts=[NoiseCount(host="google-analytics.com", count=4)],
    )
    store.save_crawl(crawl)
    home = ScreenNode(crawl_id=crawl.id, project="erp", url_template="app.test/",
                      url_example="https://app.test/", signature="sig_home", name="Home")
    settings = ScreenNode(crawl_id=crawl.id, project="erp", url_template="app.test/settings",
                          url_example="https://app.test/settings", signature="sig_set",
                          name="Settings", console_errors=["boom"])
    store.add_node(home)
    store.add_node(settings)
    store.add_edge(ScreenEdge(crawl_id=crawl.id, from_node=home.id, to_node=settings.id,
                              action=Action.CLICK, target="#settings", name="Settings",
                              outcome=EdgeOutcome.NAVIGATED))
    store.add_edge(ScreenEdge(crawl_id=crawl.id, from_node=settings.id, action=Action.CLICK,
                              target="#del", name="Delete account",
                              outcome=EdgeOutcome.DENIED_POLICY,
                              reason="destructive-name deny-list"))
    store.add_edge(ScreenEdge(crawl_id=crawl.id, from_node=settings.id, action=Action.CLICK,
                              target="#icon", name="", outcome=EdgeOutcome.SKIPPED_UNNAMED,
                              reason="unnamed control"))
    store.add_crawl_issue(CrawlIssue(crawl_id=crawl.id, project="erp", kind=IssueKind.CONSOLE,
                                     node_id=settings.id, detail="boom"))
    return store, crawl


def test_workbook_has_every_sheet(tmp_path: Path) -> None:
    seed(tmp_path)
    out = export_crawl_excel("erp", "crawl_demo", tmp_path / "r.xlsx", tmp_path)

    assert load_workbook(out).sheetnames == SHEETS


def test_summary_names_why_the_crawl_stopped(tmp_path: Path) -> None:
    """A bounded crawl that hit `max_actions` saw less of the product than one
    that emptied its frontier. A report that omits which happened invites the
    reader to assume full coverage."""
    _store, crawl = seed(tmp_path)
    assert ("Stopped because", "max_actions") in crawl_summary(crawl)


def test_refusals_sheet_lists_what_the_crawl_would_not_touch(tmp_path: Path) -> None:
    seed(tmp_path)
    out = export_crawl_excel("erp", "crawl_demo", tmp_path / "r.xlsx", tmp_path)
    ws = load_workbook(out)["Denied & Skipped"]
    rows = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]

    controls = {row[1] for row in rows}
    assert "Delete account" in controls
    assert "(unnamed)" in controls  # the icon button a human must add aria-label to
    assert len(rows) == 2  # the NAVIGATED edge is not a refusal


def test_third_party_noise_is_recorded_but_never_an_issue(tmp_path: Path) -> None:
    seed(tmp_path)
    out = export_crawl_excel("erp", "crawl_demo", tmp_path / "r.xlsx", tmp_path)
    wb = load_workbook(out)
    noise = [[c.value for c in row] for row in wb["Noise"].iter_rows(min_row=2)]
    issues = [[c.value for c in row] for row in wb["Issues"].iter_rows(min_row=2)]

    assert noise == [["google-analytics.com", 4]]
    assert all("google-analytics" not in str(row) for row in issues)


def test_screen_names_replace_opaque_ids_everywhere(tmp_path: Path) -> None:
    seed(tmp_path)
    out = export_crawl_excel("erp", "crawl_demo", tmp_path / "r.xlsx", tmp_path)
    ws = load_workbook(out)["Edges"]
    first_column = [row[0].value for row in ws.iter_rows(min_row=2)]

    assert "Home" in first_column and "Settings" in first_column
    assert not any(str(v).startswith("node_") for v in first_column)


def test_an_unknown_crawl_id_is_a_clear_error_not_an_empty_workbook(tmp_path: Path) -> None:
    seed(tmp_path)
    try:
        export_crawl_excel("erp", "crawl_nope", tmp_path / "r.xlsx", tmp_path)
    except ValueError as exc:
        assert "crawl_nope" in str(exc)
    else:  # pragma: no cover - the failure path is the point of the test
        raise AssertionError("exporting an unknown crawl silently produced a file")
