"""Every swallowed cause in `stages/explore*.py` (AT-108 + AT-114 sweep).

These are not tests that the crawl survives a failure -- `test_explore_node_
recovery.py` already covers that. They test that when it survives, it can still
SAY WHAT HAPPENED. "The cause was swallowed" is a shape, not a location: the
same `except Exception:` with an empty body appeared at three sites, and a
crawl that loses a screen, cannot screenshot one, or cannot even get back to
the base URL used to report all three identically -- as nothing at all.

Contract: qa/contracts/explore.md X11/X16.
"""

from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

import pytest
from crawl_fake import grant_crawl_approval, make_project, make_session

from autotester.browser.observe import PageObserver
from autotester.stages import explore_return
from autotester.stages.explore import run_crawl
from autotester.store.project_store import ProjectStore

# -- AT-114: a screenshot the crawler could not take ------------------------

def test_a_screenshot_failure_is_recorded_with_its_cause(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    """`session.screenshot` retries the one transient CDP race AT-036
    documents and deliberately re-raises everything else because everything
    else is real. `capture`'s bare `except Exception: return None` threw that
    distinction away: a node with no screenshot looked the same whether the
    compositor hiccuped or the browser had died."""
    project = make_project()
    session, page = make_session(tmp_path, project)

    def dead_browser(path: str, full_page: bool = False) -> None:
        raise RuntimeError("Target page, context or browser has been closed")

    monkeypatch.setattr(page, "screenshot", dead_browser)
    store = ProjectStore("demo", tmp_path)
    grant_crawl_approval(store, project)

    crawl = run_crawl(project, session, store, observer=PageObserver())

    issues = store.list_crawl_issues(crawl.id)
    shot_issues = [i for i in issues if i.kind.value == "evidence"]
    assert shot_issues, "a screenshot that never happened must leave a trace"
    assert any("browser has been closed" in i.detail for i in shot_issues)


def test_a_screenshot_failure_is_not_reported_as_a_product_bug(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    """The reason EVIDENCE is its own kind. Filing a failure of the TOOL as a
    navigation/console/network issue would inflate the product's issue count
    with the crawler's own problems -- the same dishonesty X9 forbids when it
    refuses to call third-party noise a product issue."""
    project = make_project()
    session, page = make_session(tmp_path, project)
    monkeypatch.setattr(
        page, "screenshot",
        lambda path, full_page=False: (_ for _ in ()).throw(RuntimeError("boom")))
    store = ProjectStore("demo", tmp_path)
    grant_crawl_approval(store, project)

    crawl = run_crawl(project, session, store, observer=PageObserver())

    product_kinds = {"navigation", "console", "network", "dialog"}
    leaked = [i for i in store.list_crawl_issues(crawl.id)
              if i.kind.value in product_kinds and "screenshot" in i.detail]
    assert leaked == []


# -- AT-108: a screen the crawler could not get back to ---------------------

def test_a_lost_screen_reports_why_it_was_lost(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    """`return_to` discarded its exception, so the issue it produced said only
    THAT the crawl lost a screen. On a live production crawl that is the
    difference between 'the app logged us out' and 'our own browser died'."""
    project = make_project()
    session, page = make_session(tmp_path, project)

    def broken_back(wait_until: str = "") -> None:
        raise RuntimeError("Navigation failed because page crashed")

    monkeypatch.setattr(page, "go_back", broken_back)
    store = ProjectStore("demo", tmp_path)
    grant_crawl_approval(store, project)

    crawl = run_crawl(project, session, store, observer=PageObserver())

    issues = store.list_crawl_issues(crawl.id)
    assert any("page crashed" in i.detail for i in issues), (
        "the cause reached no artifact a human reads")


def test_a_failed_recovery_says_so_instead_of_passing_silently() -> None:
    """`_recover` is the last resort -- if it fails, every screen visited after
    it is suspect. It was the most serious failure in the file and the only one
    handled with a bare `pass`."""
    def dead(url: str, wait_until: str = "") -> None:
        raise RuntimeError("browser gone")

    rt = SimpleNamespace(
        session=SimpleNamespace(goto=dead, settle=lambda timeout_ms=0: None),
        project=SimpleNamespace(base_url="https://app.test/"),
        bounds=SimpleNamespace(settle_ms=10),
        return_error="back failed",
    )
    explore_return._recover(rt)  # type: ignore[arg-type]

    assert "browser gone" in (rt.return_error or "")
    assert "back failed" in (rt.return_error or ""), "the original cause was overwritten"


# -- AT-120: the separation must hold in the NUMBERS, not just the enum -----

def test_tool_failures_are_not_counted_into_the_products_issue_total(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    """The checker broke the AT-114 headline claim by executing it: a crawl
    whose every screenshot fails reported `crawl.issues = 5` for ONE real
    product issue, because `add_issue` incremented kind-blind. The enum was
    separate; the number a human actually reads was not -- and since a failed
    screenshot previously filed no issue at all, the inflation was introduced
    by the very commit that claimed to prevent it."""
    project = make_project()
    session, page = make_session(tmp_path, project)
    monkeypatch.setattr(
        page, "screenshot",
        lambda path, full_page=False: (_ for _ in ()).throw(RuntimeError("no screenshot")))
    store = ProjectStore("demo", tmp_path)
    grant_crawl_approval(store, project)

    crawl = run_crawl(project, session, store, observer=PageObserver())

    filed = store.list_crawl_issues(crawl.id)
    product = [i for i in filed if i.kind.value != "evidence"]
    tool = [i for i in filed if i.kind.value == "evidence"]
    assert tool, "precondition: this crawl must actually fail to screenshot"
    assert crawl.issues == len(product), (
        f"headline issue count {crawl.issues} != {len(product)} real product issues")
    assert crawl.tool_failures == len(tool)


def test_the_workbook_and_the_crawl_page_keep_them_apart(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    """`crawl.issues` alone was never the whole leak -- the Issues sheet and
    the UI issue list rendered every kind together, so even a correct headline
    would have been contradicted by the rows under it."""
    from openpyxl import load_workbook

    from autotester.stages.crawl_report import export_crawl_excel
    from autotester.ui import crawl_view

    project = make_project()
    session, page = make_session(tmp_path, project)
    monkeypatch.setattr(
        page, "screenshot",
        lambda path, full_page=False: (_ for _ in ()).throw(RuntimeError("no screenshot")))
    store = ProjectStore("demo", tmp_path)
    grant_crawl_approval(store, project)
    crawl = run_crawl(project, session, store, observer=PageObserver())

    out = export_crawl_excel("demo", crawl.id, tmp_path / "r.xlsx", root=tmp_path)
    wb = load_workbook(out)
    assert "Tool failures" in wb.sheetnames
    issue_rows = list(wb["Issues"].iter_rows(min_row=2, values_only=True))
    assert not any("screenshot" in str(c) for row in issue_rows for c in row), (
        "a crawler failure reached the product's Issues sheet")

    tool = [i for i in store.list_crawl_issues(crawl.id) if i.kind.value == "evidence"]
    assert "could not screenshot" in crawl_view.tool_failures_table(tool, {}), (
        "the crawl page must still SHOW the gap in its own evidence")


# -- AT-121 / AT-122: the surfaces the AT-120 sweep missed -----------------

def _summary_crawl(**over: object):
    from autotester.schema.crawl import Crawl
    from autotester.schema.enums import CrawlStatus
    base = dict(project="demo", id="crawl_x", status=CrawlStatus.COMPLETED,
                stop_reason="frontier empty", screens=3, actions=9, edges=4,
                denied=1, issues=1, tool_failures=4)
    base.update(over)
    return Crawl(**base)  # type: ignore[arg-type]


def test_the_cli_line_reports_tool_failures_too(
    capsys: pytest.CaptureFixture[str],
) -> None:
    """AT-122 (checker-found): `cli_crawl` was the one `crawl.issues` reader
    the AT-120 commit did not touch. After that fix it no longer INFLATED the
    count -- it silently dropped tool failures instead, so a CI or headless run
    (the only report those get) was told nothing about holes in its own
    evidence. Under-reporting replaced over-reporting; both are dishonest."""
    from autotester.cli_crawl import echo_crawl_summary

    echo_crawl_summary(_summary_crawl())

    line = capsys.readouterr().out
    assert "1 issues" in line
    assert "4 tool failures" in line, "a headless run gets this line and nothing else"


def test_a_known_zero_is_printed_as_zero_not_as_unknown(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    """AT-121 (checker-found): the crawls table rendered a tool-failure count
    of 0 as the em dash it also uses for an UNKNOWN stop_reason, while the
    Issues cell beside it printed `0`. Two different meanings, one glyph, in
    adjacent columns."""
    from fastapi.testclient import TestClient

    from autotester.schema.project import Project
    from autotester.ui.app import app

    monkeypatch.setenv("AUTOTESTER_ROOT", str(tmp_path))
    store = ProjectStore("demo", tmp_path)
    store.save_project(Project(slug="demo", name="Demo", base_url="https://demo.test",
                               allowed_domains=["demo.test"]))
    store.save_crawl(_summary_crawl(id="crawl_zero", issues=0, tool_failures=0))

    html = TestClient(app).get("/projects/demo/crawls").text
    row = html[html.index("crawl_zero"):]
    row = row[:row.index("</tr>")]

    # Both count cells, and only the count cells. The trailing em dash in this
    # row is `started_at` on a crawl that never started -- a genuine unknown,
    # and the exact distinction this test exists to defend. Matching on the
    # glyph anywhere in the row would fail on that correct behaviour.
    assert row.count("<td>0</td>") == 2, (
        "issues and tool failures are both a MEASURED zero and must read alike; "
        "`or '—'` renders one of them as the sentinel for 'not known'")
