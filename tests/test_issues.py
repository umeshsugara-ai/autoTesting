"""Issue derivation and the sheet a human reads — VL5/VL6.

The sharpest test here compares the exported header against the REAL workbook
on disk when it is available, not against a list I typed. Every column-shape
fact in this file was measured on `ERP_Issues_ALL.xlsx` and
`ERP_Issues_Trainers.xlsx` in 2026-09-08 (`.work/track-a-corpus-facts.md`), and
four of them contradicted the plan.

Contract: qa/contracts/video-learning.md VL5/VL6.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from autotester.schema.analysis import AnalysedIssue, VideoAnalysis
from autotester.schema.enums import (
    Confidence,
    IssueCategory,
    IssueOrigin,
    Severity,
    SourceKind,
)
from autotester.schema.project import Source
from autotester.stages.issues import (
    ISSUE_COLUMNS,
    at_mmss,
    derive_issues,
    export_issues_excel,
    issue_row,
)

HUMAN_SHEET = Path("C:/Users/Lenovo/Videos/Screen Recordings/ERP_Issues_ALL.xlsx")


def a_source(**kw) -> Source:
    return Source(project="erp", kind=SourceKind.VIDEO, path="/x/erp1.mp4",
                  label=kw.pop("label", "erp1.mp4 (Divya Kamboj, trainer pipeline)"),
                  recorded_on=kw.pop("recorded_on", "2026-08-18"), sha256="d", **kw)


def an_issue(**kw) -> AnalysedIssue:
    return AnalysedIssue(
        t_start=kw.pop("t_start", 93.4), screen=kw.pop("screen", "Trainers"),
        category=kw.pop("category", IssueCategory.FEATURE_GAP),
        severity=kw.pop("severity", Severity.S1),
        title=kw.pop("title", "Trainer cannot be moved past Shortlisted"),
        what_is_wrong=kw.pop("what_is_wrong", "the Move button does nothing"), **kw)


def an_analysis(*issues: AnalysedIssue) -> VideoAnalysis:
    return VideoAnalysis(source_id="src_1", issues=list(issues))


# -- the header is theirs, not ours ----------------------------------------

@pytest.mark.skipif(not HUMAN_SHEET.exists(), reason="the real corpus is not on this host")
def test_the_header_matches_the_real_human_sheet() -> None:
    """The whole point of T-136 is that a tester can put our sheet beside their
    own and read across, so the shape has to match THEIRS.

    Compared against the workbook on disk rather than a list I typed — a
    hand-copied header is a claim about a file, and this reads the file."""
    workbook = load_workbook(HUMAN_SHEET, read_only=True)
    theirs = [cell.value for cell in next(workbook["All issues"].iter_rows(max_row=1))]
    workbook.close()

    assert theirs == ISSUE_COLUMNS


def test_the_header_is_thirteen_columns_in_order() -> None:
    """The offline half, so the shape is still pinned on a host without the
    corpus. If this and the sheet ever disagree, the sheet wins."""
    assert len(ISSUE_COLUMNS) == 13
    assert ISSUE_COLUMNS[0] == "ID"
    assert ISSUE_COLUMNS[-1] == "Evidence"
    assert "Said verbatim" in ISSUE_COLUMNS


def test_the_workbook_writes_that_header_and_one_row_per_issue(tmp_path: Path) -> None:
    out = export_issues_excel(
        derive_issues(an_analysis(an_issue(), an_issue(t_start=10.0, screen="Home")),
                      a_source(), "erp"),
        tmp_path / "issues.xlsx")

    sheet = load_workbook(out)["All issues"]
    assert [c.value for c in next(sheet.iter_rows(max_row=1))] == ISSUE_COLUMNS
    assert sheet.max_row == 3


# -- the two fields a scorer compares on ------------------------------------

@pytest.mark.parametrize("seconds,expected", [
    (0.0, "00:00"), (9.4, "00:09"), (29.6, "00:30"), (93.4, "01:33"), (600.0, "10:00"),
])
def test_times_are_written_the_way_the_human_sheet_writes_them(
    seconds: float, expected: str,
) -> None:
    """Measured on the real workbook: the `At` column holds STRINGS like
    `00:29`, not numbers. A scorer comparing a float to that cell matches
    nothing, silently — one of four corpus facts that would each have produced
    a recall of zero."""
    assert at_mmss(seconds) == expected


def test_severity_is_written_in_their_words_not_ours() -> None:
    """A tester reading `S2` has to translate; a tester reading `Medium` does
    not, and this sheet is for them."""
    row = issue_row(derive_issues(an_analysis(an_issue(severity=Severity.S1)),
                                  a_source(), "erp")[0])

    assert row[ISSUE_COLUMNS.index("Severity")] == "High"


def test_the_recording_column_carries_the_label_a_human_would_recognise() -> None:
    """The human sheet's clip column reads
    `erp1.mp4 (Divya Kamboj, trainer pipeline)` — a file plus who and what.
    Writing a bare source id there would make the two sheets unjoinable."""
    row = issue_row(derive_issues(an_analysis(an_issue()), a_source(), "erp")[0])

    assert row[ISSUE_COLUMNS.index("Recording")].startswith("erp1.mp4")


# -- how we know: derived from evidence, not from the model's claim ---------

@pytest.mark.parametrize("narration,on_screen,expected", [
    ("this should say Save", "Certified 0", IssueOrigin.SPOKEN_AND_SCREEN),
    ("this should say Save", None, IssueOrigin.SPOKEN),
    (None, "Certified 0", IssueOrigin.SCREEN),
    (None, None, IssueOrigin.MODEL_DETECTED),
])
def test_how_we_know_follows_the_evidence_actually_present(
    narration: str | None, on_screen: str | None, expected: IssueOrigin,
) -> None:
    """A model asked "how do you know" answers plausibly every time. The fields
    either hold a quote and on-screen text or they do not, so this reads them
    instead of asking."""
    issue = derive_issues(
        an_analysis(an_issue(narration=narration, on_screen_text=on_screen)),
        a_source(), "erp")[0]

    assert issue.how_we_know is expected


def test_a_testers_words_reach_the_said_verbatim_column() -> None:
    """`narration` on an observation becomes `said_verbatim` on the artifact.
    It is the highest-value column in the sheet — a tester's own sentence — and
    it survives the rename or it is lost."""
    row = issue_row(derive_issues(
        an_analysis(an_issue(narration="this should say Save")), a_source(), "erp")[0])

    assert row[ISSUE_COLUMNS.index("Said verbatim")] == "this should say Save"


# -- provenance and idempotence --------------------------------------------

def test_re_deriving_the_same_analysis_gives_the_same_ids() -> None:
    """`Issue.id` is content-addressed, which is what lets `add_issue` stay
    idempotent — a re-run must not double the ledger."""
    analysis, source = an_analysis(an_issue()), a_source()

    first = derive_issues(analysis, source, "erp")
    second = derive_issues(analysis, source, "erp")

    assert [i.id for i in first] == [i.id for i in second]
    assert first[0].id


def test_every_issue_points_back_at_its_recording_and_second() -> None:
    """The id comes from the SOURCE, not from the analysis's own `source_id`
    field — my first version asserted the latter and failed. The source is what
    a human can open; the analysis is a derived artifact about it."""
    source = a_source()
    issue = derive_issues(an_analysis(an_issue(t_start=93.4)), source, "erp")[0]

    assert issue.source_id == source.id
    assert issue.at_s == 93.4
    assert issue.recorded_on == "2026-08-18"


def test_model_agreement_survives_into_the_artifact() -> None:
    """Two models describing the same fault independently is the strongest
    signal this pipeline produces without a human, and it must reach the row."""
    issue = derive_issues(
        an_analysis(an_issue(models_agreeing=2, model_labels=["a", "b"],
                             confidence=Confidence.HIGH)),
        a_source(), "erp")[0]

    assert issue.models_agreeing == 2
    assert issue.confidence is Confidence.HIGH


def test_an_analysis_with_no_issues_writes_a_header_only_sheet(tmp_path: Path) -> None:
    """An empty sheet is a real answer — "we found nothing" — and must not be a
    crash or a missing file."""
    out = export_issues_excel(derive_issues(an_analysis(), a_source(), "erp"),
                              tmp_path / "empty.xlsx")

    sheet = load_workbook(out)["All issues"]
    assert sheet.max_row == 1
    assert [c.value for c in next(sheet.iter_rows(max_row=1))] == ISSUE_COLUMNS


def test_the_At_CELL_carries_MM_SS_not_a_number(tmp_path: Path) -> None:
    """Sabotaging `issue_row` to write a raw float came back INCONCLUSIVE,
    because the parametrized test above exercises `at_mmss` DIRECTLY and never
    the row. The helper being right does not make the sheet right — that is
    this session's whole recurring failure, arriving in the file about matching
    a human's sheet.

    So this reads the written workbook, which is what a tester opens."""
    out = export_issues_excel(
        derive_issues(an_analysis(an_issue(t_start=93.4)), a_source(), "erp"),
        tmp_path / "issues.xlsx")

    sheet = load_workbook(out)["All issues"]
    row = next(sheet.iter_rows(min_row=2, max_row=2, values_only=True))
    at_cell = row[ISSUE_COLUMNS.index("At")]

    assert at_cell == "01:33", f"the At cell reads {at_cell!r}, not MM:SS"
    assert isinstance(at_cell, str)
