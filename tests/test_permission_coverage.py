"""Coverage is measured against what the ACCOUNT may do (coverage.md V9, T-171, AT-663).

Umesh, verbatim: "jo account mai dunga usme jitni permission hogi uthi tho testing ho hi
jaani chaiyee." A figure whose denominator is "the controls a bounded crawl happened to
find" can read 100 % while every control the supplied account is entitled to press sits
untouched. These tests pin the three things V9 forbids: choosing the denominator by what
was seen, dropping a blocked control out of the denominator, and presenting a
screens-reached figure as permission coverage.

Lives beside `test_crawl_coverage.py` rather than inside it: that file is 251 lines against
C2's 300-line cap, so V9 could not be added to it without breaking the cap this repo
enforces. Same subject, one file per criterion cluster — the shape `test_crawl_coverage_
bounds.py` and `test_coverage_wiring.py` already set.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from autotester.schema.crawl import Crawl, CrawlBounds, PermittedControl
from autotester.schema.enums import CrawlStatus, EdgeOutcome, NodeStatus, WritePolicy
from autotester.schema.project import Project
from autotester.schema.screen_graph import ElementRef, ScreenEdge, ScreenNode
from autotester.stages.crawl_coverage import REASONS, compute_coverage
from autotester.stages.crawl_report import coverage_figure, coverage_rows, export_crawl_excel
from autotester.store.project_store import ProjectStore
from autotester.ui.app import app

TEMPLATE = "/app"
NODE_TEMPLATE = "app.test/app"
"""Deliberately host-FUL, the shape `screen_identity.node_from` really stores, while the
declared surface above is host-less: a project's declaration of what its account may do
must not be bound to the environment a crawl happened to run against."""


def _node(*selectors: str, status: NodeStatus = NodeStatus.EXPLORED) -> ScreenNode:
    return ScreenNode(crawl_id="c", project="demo", url_template=NODE_TEMPLATE,
                      url_example="https://app.test/app", signature="s", status=status,
                      elements=[ElementRef(role="button", name=s, selector=s)
                                for s in selectors])


def _edge(node: ScreenNode, target: str, outcome: EdgeOutcome, **extra: object) -> ScreenEdge:
    return ScreenEdge(crawl_id="c", from_node=node.id, target=target, action="click",
                      outcome=outcome, **extra)


def _permits(*selectors: str, destructive: bool = False) -> list[PermittedControl]:
    return [PermittedControl(url_template=TEMPLATE, selector=s, name=s.lstrip("#"),
                             destructive=destructive) for s in selectors]


def _compute(nodes: list[ScreenNode], edges: list[ScreenEdge], *,
             permitted: list[PermittedControl] | None = None,
             policy: WritePolicy = WritePolicy.READ_ONLY,
             status: CrawlStatus = CrawlStatus.COMPLETED, bound: str | None = None):
    return compute_coverage(status=status, bound=bound, bounds=CrawlBounds(), nodes=nodes,
                            edges=edges, permitted=permitted, policy=policy)


# -- the denominator ----------------------------------------------------------------------

def test_the_denominator_is_the_permitted_surface_not_the_screens_reached() -> None:
    """V9's core: the crawl found and exercised ONE control and its frontier drained; the
    account permits four. A screens-reached denominator calls that 100 %."""
    node = _node("#a")

    cov = _compute([node], [_edge(node, "#a", EdgeOutcome.SAME_SCREEN)],
                   permitted=_permits("#a", "#b", "#c", "#d"))

    assert cov.controls_exercised == cov.controls_discovered == 1  # what it SAW
    assert cov.denominator_basis == "permitted"
    assert (cov.permitted_exercised, cov.permitted_total) == (1, 4)
    assert cov.percent == 25, "the figure must be a fraction of the permitted surface"


def test_a_permitted_control_no_entry_point_reached_is_listed_not_dropped() -> None:
    node = _node("#a")

    cov = _compute([node], [_edge(node, "#a", EdgeOutcome.SAME_SCREEN)],
                   permitted=_permits("#a", "#never"))

    holes = {h.selector: h.reason for h in cov.permitted_holes}
    assert holes == {"#never": "not_reached"}, "not reached is a REASON, not an absence of one"
    assert cov.permitted_total == 2, "the denominator still includes it"


def test_a_blocked_control_is_reported_blocked_and_never_counts_as_covered() -> None:
    """The hard rule: a control the run refused is a hole with the refusal as its reason,
    and it stays in the denominator."""
    node = _node("#a", "#pay")
    edges = [_edge(node, "#a", EdgeOutcome.SAME_SCREEN),
             _edge(node, "#pay", EdgeOutcome.DENIED_POLICY, reason="deny-list: pay")]

    cov = _compute([node], edges, permitted=_permits("#a", "#pay"))

    blocked = [(h.selector, h.reason) for h in cov.permitted_holes]
    assert blocked == [("#pay", "policy:deny-list: pay")]
    assert (cov.permitted_exercised, cov.permitted_total) == (1, 2)
    assert cov.percent == 50


def test_a_declared_destructive_control_names_the_write_policy_that_stood_in_the_way() -> None:
    """It was never reached, but `not_reached` would hide WHY. Under ALLOW_WRITES the same
    control is an ordinary miss — the tier is what makes the surface reachable."""
    node = _node("#a")
    edges = [_edge(node, "#a", EdgeOutcome.SAME_SCREEN)]
    permitted = _permits("#a") + _permits("#delete", destructive=True)

    blocked = _compute([node], edges, permitted=permitted, policy=WritePolicy.READ_ONLY)
    opened = _compute([node], edges, permitted=permitted, policy=WritePolicy.ALLOW_WRITES)

    assert [h.reason for h in blocked.permitted_holes] == ["policy:destructive under read_only"]
    assert [h.reason for h in opened.permitted_holes] == ["not_reached"]


def test_every_permitted_hole_carries_one_reason_from_the_closed_set() -> None:
    node = _node("#a", "#un", status=NodeStatus.EXPLORED)
    edges = [_edge(node, "#a", EdgeOutcome.SAME_SCREEN),
             _edge(node, "#un", EdgeOutcome.SKIPPED_UNNAMED)]

    cov = _compute([node], edges, permitted=_permits("#a", "#un", "#gone")
                   + _permits("#burn", destructive=True))

    assert len(cov.permitted_holes) == 3
    for hole in cov.permitted_holes:
        assert any(hole.reason == r or (r.endswith(":") and hole.reason.startswith(r))
                   for r in REASONS), hole.reason
    assert (cov.permitted_exercised + len(cov.permitted_holes)) == cov.permitted_total


def test_a_fully_exercised_permitted_surface_may_read_100() -> None:
    """The guard must be able to say yes, or the figure is not a measure."""
    node = _node("#a")

    cov = _compute([node], [_edge(node, "#a", EdgeOutcome.SAME_SCREEN)], permitted=_permits("#a"))

    assert cov.percent == 100 and not cov.permitted_holes


def test_a_bounded_crawl_over_a_permitted_surface_still_never_reads_100() -> None:
    """V7(d) survives the re-base: an incomplete crawl cannot claim the account's whole surface."""
    node = _node("#a")

    cov = _compute([node], [_edge(node, "#a", EdgeOutcome.SAME_SCREEN)], permitted=_permits("#a"),
                   status=CrawlStatus.STOPPED_BOUND, bound="max_actions")

    assert cov.percent < 100


# -- the honest label when the permitted surface is unknown --------------------------------

def test_without_a_declared_surface_the_report_says_its_denominator_is_screens_reached() -> None:
    node = _node("#a")

    cov = _compute([node], [_edge(node, "#a", EdgeOutcome.SAME_SCREEN)])
    figure = coverage_figure(Crawl(project="demo", coverage=cov))

    assert cov.denominator_basis == "screens-reached"
    assert cov.permitted_total is None and not cov.permitted_holes
    assert "NOT permission coverage" in figure, figure


def test_with_a_declared_surface_the_figure_names_the_account_role() -> None:
    node = _node("#a")

    cov = _compute([node], [_edge(node, "#a", EdgeOutcome.SAME_SCREEN)],
                   permitted=_permits("#a", "#b"))
    crawl = Crawl(project="demo", coverage=cov)
    rows = dict(coverage_rows(crawl))

    assert "the account's role permits" in coverage_figure(crawl)
    assert "NOT permission coverage" not in coverage_figure(crawl)
    assert rows["Coverage denominator"] == "permitted"
    assert rows["Permitted but not exercised, by reason"] == "not_reached: 1"


# -- shown: the crawl page and the workbook ------------------------------------------------

@pytest.fixture
def stored(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> tuple[Path, Crawl]:
    monkeypatch.setenv("AUTOTESTER_ROOT", str(tmp_path))
    store = ProjectStore("demo", tmp_path)
    store.save_project(Project(slug="demo", name="Demo", base_url="https://app.test/",
                               allowed_domains=["app.test"], permitted_surface=_permits("#a")))
    node = _node("#a")
    cov = _compute([node], [_edge(node, "#a", EdgeOutcome.SAME_SCREEN)],
                   permitted=_permits("#a", "#b") + _permits("#pay", destructive=True))
    crawl = Crawl(id="crawl_perm", project="demo", status=CrawlStatus.COMPLETED, screens=1,
                  actions=1, stop_reason="frontier empty", coverage=cov)
    store.save_crawl(crawl)
    return tmp_path, crawl


def test_the_crawl_page_shows_the_permitted_denominator_and_every_hole(
    stored: tuple[Path, Crawl],
) -> None:
    _root, crawl = stored

    text = TestClient(app).get(f"/projects/demo/crawls/{crawl.id}").text

    assert "the account&#x27;s role permits" in text or "the account's role permits" in text
    assert "permitted but not exercised" in text
    assert "not_reached: 1" in text and "policy:destructive under read_only: 1" in text
    assert "Coverage denominator" in text


def test_the_workbook_lists_every_permitted_control_not_exercised(
    stored: tuple[Path, Crawl], tmp_path: Path,
) -> None:
    root, crawl = stored

    wb = load_workbook(export_crawl_excel("demo", crawl.id, tmp_path / "p.xlsx", root))
    summary = {r[0]: r[1] for r in wb["Summary"].iter_rows(values_only=True)}
    unreached = [r for r in list(wb["Unreached"].iter_rows(values_only=True))[1:]
                 if str(r[1]).startswith("permitted, not exercised")]

    assert summary["Coverage denominator"] == "permitted"
    assert {r[2] for r in unreached} == {"#b", "#pay"}
    assert {r[3] for r in unreached} == {"not_reached", "policy:destructive under read_only"}


def test_the_crawl_reads_the_projects_declared_surface_end_to_end(tmp_path: Path) -> None:
    """The wire (V6's device, applied to V9): `of_run` must take the denominator from the
    PROJECT, not from a caller's argument, or a declared surface never reaches a real crawl."""
    from crawl_fake import crawl_it, make_project

    project = make_project().model_copy(
        update={"permitted_surface": _permits("#nothing-like-this")})
    crawl, _store, _page = crawl_it(tmp_path, project=project)

    assert crawl.coverage is not None
    assert crawl.coverage.denominator_basis == "permitted"
    assert crawl.coverage.permitted_total == 1
    assert [h.reason for h in crawl.coverage.permitted_holes] == ["not_reached"]
    assert crawl.coverage.percent == 0


# -- PS2 / D-040: destructive actions are ordered last -------------------------------------

def test_a_destructive_control_is_attempted_after_every_ordinary_one_on_its_screen() -> None:
    """D-040 verbatim. A stable partition: the order inside each half is untouched, so the
    only difference is where the `Delete`/`Pay` control sits."""
    from autotester.schema.crawl import SafetyPolicy
    from autotester.stages.explore_safety import destructive_last

    elements = [ElementRef(role="button", name=n, selector=f"#{i}")
                for i, n in enumerate(["Delete account", "Open", "Pay now", "Save", "Edit"])]

    ordered = destructive_last(elements, SafetyPolicy())

    assert [el.name for el in ordered] == ["Open", "Save", "Edit", "Delete account", "Pay now"]


def test_the_crawl_really_tries_the_destructive_control_last(tmp_path: Path) -> None:
    """The wire, not the helper: the edges a real (fake-site) crawl RECORDED for the screen
    that carries `Delete account` must put it after every other control on that screen."""
    from crawl_fake import crawl_it

    crawl, store, _page = crawl_it(tmp_path)
    settings = next(n for n in store.list_nodes(crawl.id) if n.url_template.endswith("/settings"))
    order = [e.target for e in store.list_edges(crawl.id) if e.from_node == settings.id]

    assert "button.del" in order, "the destructive control must still be accounted for"
    assert order.index("button.del") == len(order) - 1, order


def test_every_destructive_press_comes_after_every_non_destructive_one_across_the_crawl(
        tmp_path: Path) -> None:
    """PS2 [D-040 verbatim], CRAWL-GLOBAL: not per screen. The checker's reproduction is a
    multi-screen crawl under ALLOW_WRITES in which `button.del` on /settings was pressed
    before `button.edit` on a screen the crawl visits later."""
    from crawl_fake import crawl_it, make_project

    from autotester.schema.crawl import SafetyPolicy
    from autotester.stages.explore_safety import is_destructive

    policy = SafetyPolicy(write_policy=WritePolicy.ALLOW_WRITES)
    crawl, store, _page = crawl_it(tmp_path, project=make_project(WritePolicy.ALLOW_WRITES),
                                   policy=policy)
    pressed = {EdgeOutcome.NAVIGATED, EdgeOutcome.SAME_SCREEN}
    seq = [(i, e) for i, e in enumerate(store.list_edges(crawl.id)) if e.outcome in pressed]
    kind = [is_destructive(ElementRef(role="button", name=e.name or "", selector=e.target),
                           policy) for _i, e in seq]

    assert any(kind), "the fixture must really press a destructive control"
    assert not all(kind), "the fixture must press ordinary controls too"
    first_destructive = kind.index(True)
    assert not any(not k for k in kind[first_destructive:]), [
        (e.target, e.outcome.value, k) for (_i, e), k in zip(seq, kind, strict=True)]
    assert len({e.from_node for _i, e in seq}) > 1, "the crawl must span several screens"
