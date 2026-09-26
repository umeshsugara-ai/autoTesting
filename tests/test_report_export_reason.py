"""Contract: qa/contracts/report-export.md RE6 (D-045). AT-582.

Both exports must show, for a FAIL or INCONCLUSIVE verdict, each failure's
criterion, reason and fix_hint, plus the case's own steps as the repro --
verbatim from the stored Verdict/Case (RE1: nothing recomputed).
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from autotester.schema.case import Case
from autotester.schema.enums import Action, CaseClass, CaseKind, Outcome, Result
from autotester.schema.flowspec import Step
from autotester.schema.project import Project
from autotester.schema.run import RawResult, Run
from autotester.schema.verdict import Failure, Verdict
from autotester.stages import report_export
from autotester.store import ProjectStore

RUN_ID = "run-reason123"


def _seed(
    tmp_path: Path,
    *,
    outcome: Outcome,
    result: Result,
    failures: list[Failure] | None = None,
    error: str | None = None,
) -> ProjectStore:
    store = ProjectStore("demo", tmp_path)
    store.save_project(
        Project(slug="demo", name="Demo", base_url="https://demo.test",
                allowed_domains=["demo.test"])
    )
    case = Case(
        project="demo", flow_id="flow-login", kind=CaseKind.BEST, case_class=CaseClass.HAPPY,
        title="Login works",
        steps=[
            Step(order=1, action=Action.NAVIGATE, target="/login"),
            Step(order=2, action=Action.FILL, target="role=textbox[name=Email]",
                 value="a@b.test"),
            Step(order=3, action=Action.CLICK, target="role=button[name=Sign in]"),
        ],
    )
    store.add_case(case)
    store.save_run(Run(id=RUN_ID, project="demo", case_ids=[case.id]))
    raw = RawResult(case_id=case.id, outcome=outcome, duration_s=1.0, error=error)
    store.save_result(RUN_ID, raw)
    store.save_verdict(RUN_ID, Verdict(
        run_id=RUN_ID, case_id=case.id, result=result, criteria_met=0, criteria_total=1,
        scoreboard="Criteria 0/1 met.", grader_provider="gemini",
        failures=failures or [],
    ))
    return store


def test_export_excel_shows_failure_reason_and_fix_hint_for_a_fail(tmp_path: Path) -> None:
    _seed(
        tmp_path, outcome=Outcome.COMPLETED, result=Result.FAIL,
        failures=[Failure(criterion_id="c1", reason="login form still visible",
                           fix_hint="check the redirect wait")],
    )

    out = report_export.export_excel("demo", RUN_ID, tmp_path / "out.xlsx", tmp_path)

    wb = load_workbook(out)
    rows = list(wb.active.iter_rows(values_only=True))
    header = rows[0]
    assert "Failures" in header
    assert "Repro steps" in header
    row = dict(zip(header, rows[1], strict=True))
    assert "c1" in row["Failures"]
    assert "login form still visible" in row["Failures"]
    assert "check the redirect wait" in row["Failures"]
    assert "navigate /login" in row["Repro steps"]
    assert "fill role=textbox[name=Email] = a@b.test" in row["Repro steps"]


def test_export_excel_existing_columns_are_unchanged(tmp_path: Path) -> None:
    _seed(tmp_path, outcome=Outcome.COMPLETED, result=Result.PASS)

    out = report_export.export_excel("demo", RUN_ID, tmp_path / "out.xlsx", tmp_path)

    wb = load_workbook(out)
    rows = list(wb.active.iter_rows(values_only=True))
    assert rows[0][:9] == (
        "Case", "Kind", "Class", "Outcome", "Result", "Criteria met",
        "Duration (s)", "Grader", "Notes",
    )
    assert rows[1][0] == "Login works"
    assert rows[1][4] == "PASS"


def test_export_excel_leaves_failure_columns_blank_for_a_pass(tmp_path: Path) -> None:
    _seed(tmp_path, outcome=Outcome.COMPLETED, result=Result.PASS)

    out = report_export.export_excel("demo", RUN_ID, tmp_path / "out.xlsx", tmp_path)

    wb = load_workbook(out)
    rows = list(wb.active.iter_rows(values_only=True))
    header = rows[0]
    row = dict(zip(header, rows[1], strict=True))
    assert not row["Failures"]
    assert not row["Repro steps"]


def test_export_excel_shows_repro_steps_for_inconclusive_even_with_no_failures(
    tmp_path: Path,
) -> None:
    _seed(tmp_path, outcome=Outcome.ERRORED, result=Result.INCONCLUSIVE, error="boom")

    out = report_export.export_excel("demo", RUN_ID, tmp_path / "out.xlsx", tmp_path)

    wb = load_workbook(out)
    rows = list(wb.active.iter_rows(values_only=True))
    header = rows[0]
    row = dict(zip(header, rows[1], strict=True))
    assert "navigate /login" in row["Repro steps"]
    assert not row["Failures"]  # openpyxl round-trips "" as None


def test_export_html_shows_failure_criterion_reason_and_fix_hint(tmp_path: Path) -> None:
    _seed(
        tmp_path, outcome=Outcome.COMPLETED, result=Result.FAIL,
        failures=[Failure(criterion_id="c1", reason="login form still visible",
                           fix_hint="check the redirect wait")],
    )

    out = report_export.export_html("demo", RUN_ID, tmp_path / "out.html", tmp_path)

    html = out.read_text(encoding="utf-8")
    assert "c1" in html
    assert "login form still visible" in html
    assert "check the redirect wait" in html


def test_export_html_shows_repro_steps(tmp_path: Path) -> None:
    _seed(
        tmp_path, outcome=Outcome.COMPLETED, result=Result.FAIL,
        failures=[Failure(criterion_id="c1", reason="x", fix_hint=None)],
    )

    out = report_export.export_html("demo", RUN_ID, tmp_path / "out.html", tmp_path)

    html = out.read_text(encoding="utf-8")
    assert "navigate /login" in html
    assert "fill role=textbox[name=Email] = a@b.test" in html


def test_export_html_omits_detail_block_for_a_pass(tmp_path: Path) -> None:
    _seed(tmp_path, outcome=Outcome.COMPLETED, result=Result.PASS)

    out = report_export.export_html("demo", RUN_ID, tmp_path / "out.html", tmp_path)

    html = out.read_text(encoding="utf-8")
    assert "class='detail'" not in html
