"""Contract: qa/contracts/report-export.md RE5. AT-594.

T-183 made `Case.steps` (including a FILL step's `Step.value`) the first thing
rendered into an exported artifact (`report_export._repro_steps`). RE5 already
promises no secret ever reaches an export; this pins it for repro steps
specifically:

- a FILL step carrying a `{{SECRET:KEY}}` placeholder exports as the
  placeholder, never a resolved value (the placeholder is never substituted
  in a stored `Case` -- substitution happens only at `page.fill()` time);
- a step whose value is somehow a declared raw secret is redacted by the
  existing `Redactor`/`SecretStore` path (core.redact / browser.secrets) in
  both export formats, the same last-stop guarantee every other artifact gets.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from autotester.schema.case import Case
from autotester.schema.enums import Action, CaseClass, CaseKind, Outcome, Result
from autotester.schema.flowspec import Step
from autotester.schema.project import Project, SecretRef
from autotester.schema.run import RawResult, Run
from autotester.schema.verdict import Verdict
from autotester.stages import report_export
from autotester.store import ProjectStore

RUN_ID = "run-secret123"
SECRET_KEY = "DEMO_LOGIN_PASSWORD"
SECRET_VALUE = "hunter2-super-secret"  # a fake test value, never a real credential


def _seed(tmp_path: Path, *, fill_value: str) -> ProjectStore:
    """A project that DECLARES `SECRET_KEY`, with a real `.env` value for it
    (AT-594 needs a loadable SecretStore to exercise the redaction path), and
    one FAIL case whose FILL step carries `fill_value` verbatim -- FAIL so RE6
    detail (and therefore the repro steps) is shown in both exports."""
    store = ProjectStore("demo", tmp_path)
    store.save_project(
        Project(
            slug="demo", name="Demo", base_url="https://demo.test",
            allowed_domains=["demo.test"],
            secrets=[SecretRef(key=SECRET_KEY, domains=["demo.test"])],
        )
    )
    (tmp_path / ".env").write_text(f"{SECRET_KEY}={SECRET_VALUE}\n", encoding="utf-8")
    case = Case(
        project="demo", flow_id="flow-login", kind=CaseKind.BEST, case_class=CaseClass.HAPPY,
        title="Login works",
        steps=[
            Step(order=1, action=Action.NAVIGATE, target="/login"),
            Step(order=2, action=Action.FILL, target="role=textbox[name=Password]",
                 value=fill_value),
        ],
    )
    store.add_case(case)
    store.save_run(Run(id=RUN_ID, project="demo", case_ids=[case.id]))
    store.save_result(RUN_ID, RawResult(case_id=case.id, outcome=Outcome.COMPLETED, duration_s=1.0))
    store.save_verdict(RUN_ID, Verdict(
        run_id=RUN_ID, case_id=case.id, result=Result.FAIL, criteria_met=0, criteria_total=1,
        scoreboard="Criteria 0/1 met.", grader_provider="gemini",
    ))
    return store


def _excel_repro_text(tmp_path: Path, out_name: str) -> str:
    out = report_export.export_excel("demo", RUN_ID, tmp_path / out_name, tmp_path)
    wb = load_workbook(out)
    rows = list(wb.active.iter_rows(values_only=True))
    header = rows[0]
    row = dict(zip(header, rows[1], strict=True))
    return row["Repro steps"] or ""


def test_export_excel_repro_steps_show_the_placeholder_never_the_secret_value(
    tmp_path: Path,
) -> None:
    _seed(tmp_path, fill_value=f"{{{{SECRET:{SECRET_KEY}}}}}")

    repro = _excel_repro_text(tmp_path, "out.xlsx")

    assert f"{{{{SECRET:{SECRET_KEY}}}}}" in repro
    assert SECRET_VALUE not in repro


def test_export_html_repro_steps_show_the_placeholder_never_the_secret_value(
    tmp_path: Path,
) -> None:
    _seed(tmp_path, fill_value=f"{{{{SECRET:{SECRET_KEY}}}}}")

    out = report_export.export_html("demo", RUN_ID, tmp_path / "out.html", tmp_path)

    html = out.read_text(encoding="utf-8")
    assert f"{{{{SECRET:{SECRET_KEY}}}}}" in html
    assert SECRET_VALUE not in html


def test_export_excel_redacts_a_raw_secret_value_in_a_step(tmp_path: Path) -> None:
    _seed(tmp_path, fill_value=SECRET_VALUE)

    repro = _excel_repro_text(tmp_path, "out.xlsx")

    assert SECRET_VALUE not in repro
    assert f"[REDACTED]:{SECRET_KEY}" in repro


def test_export_html_redacts_a_raw_secret_value_in_a_step(tmp_path: Path) -> None:
    _seed(tmp_path, fill_value=SECRET_VALUE)

    out = report_export.export_html("demo", RUN_ID, tmp_path / "out.html", tmp_path)

    html = out.read_text(encoding="utf-8")
    assert SECRET_VALUE not in html
    assert f"[REDACTED]:{SECRET_KEY}" in html


def test_export_excel_redacted_value_never_reaches_any_cell_in_the_workbook(
    tmp_path: Path,
) -> None:
    """Belt-and-braces: the raw value must not leak into ANY column (e.g. via
    a future column reusing case data), not only the Repro steps cell."""
    _seed(tmp_path, fill_value=SECRET_VALUE)

    out = report_export.export_excel("demo", RUN_ID, tmp_path / "out.xlsx", tmp_path)
    wb = load_workbook(out)
    all_text = "\n".join(
        str(cell) for row in wb.active.iter_rows(values_only=True) for cell in row if cell
    )

    assert SECRET_VALUE not in all_text
