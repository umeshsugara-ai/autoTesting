"""T-190 advisory UX track, part 2: exports (PU7), condition alignment (PU8), budget and
honesty. Contract: qa/contracts/persona-ux-advisory.md. Fixtures live in test_user_persona.py."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook
from test_user_persona import RUN, SECRET, _draft, _mock, _persona, _seed, _status

from autotester.providers.base import ProviderError
from autotester.providers.mock import MockProvider
from autotester.schema.enums import CaseClass, Outcome
from autotester.schema.ux_report import UXCaseStatus, UXJudgment
from autotester.schema.verdict import Result, Verdict
from autotester.stages.report_export import export_excel, export_html
from autotester.stages.ux_advisory import alignment_refusal, run_ux_advisory
from autotester.store.project_store import ProjectStore


# -- PU7: findings render as their own section ----------------------------------------------
def _fail_with_ux(tmp_path: Path) -> ProjectStore:
    store = _seed(tmp_path)
    case = store.list_cases()[0]
    store.save_verdict(RUN, Verdict(run_id=RUN, case_id=case.id, result=Result.FAIL,
                                    criteria_total=1, grader_provider="mock", failures=[
        {"criterion_id": "c1", "reason": "totals wrong", "evidence_refs": ["0.png"]}]))
    run_ux_advisory(store, RUN, _mock(UXJudgment(findings=[_draft(text=f"Jargon {SECRET}")])))
    return store


def test_persona_findings_are_a_sibling_html_section_and_a_separate_sheet(tmp_path: Path) -> None:
    store = _fail_with_ux(tmp_path)
    html = export_html("demo", RUN, tmp_path / "r.html", tmp_path).read_text(encoding="utf-8")
    assert "<h3>Failures</h3>" in html and "UX findings (advisory" in html
    detail = html[html.index("<div class='detail'>"):html.index("<div class='ux-findings'")]
    assert "Jargon" not in detail and detail.count("<div") == detail.count("</div>")
    assert SECRET not in html
    wb = load_workbook(export_excel("demo", RUN, tmp_path / "r.xlsx", tmp_path))
    main, ux = wb["Run report"], wb["UX findings (advisory)"]
    assert not any("Jargon" in str(c.value) for row in main.iter_rows() for c in row)
    assert any("Jargon" in str(c.value) for row in ux.iter_rows() for c in row)
    assert store.load_ux_report(RUN) is not None


def test_persona_no_findings_means_no_section_and_a_broken_report_is_not_zero_findings(
    tmp_path: Path,
) -> None:
    store = _seed(tmp_path)
    run_ux_advisory(store, RUN, _mock(UXJudgment(findings=[])))
    html = export_html("demo", RUN, tmp_path / "a.html", tmp_path).read_text(encoding="utf-8")
    assert "ux-findings" not in html
    assert "UX findings (advisory)" not in load_workbook(
        export_excel("demo", RUN, tmp_path / "a.xlsx", tmp_path)).sheetnames
    store.paths.run_ux_report(RUN).write_text("{not json", encoding="utf-8")
    html = export_html("demo", RUN, tmp_path / "b.html", tmp_path).read_text(encoding="utf-8")
    assert "UX report unavailable" in html
    with pytest.raises(ValueError):
        store.load_ux_report(RUN)


# -- PU8: the claim must match what ran ---------------------------------------------------
@pytest.mark.parametrize(("device", "locale", "case_class", "outcome", "refused"), [
    ("mobile", "en-US", CaseClass.HAPPY, Outcome.COMPLETED, True),  # claims a viewport not run
    ("desktop", "en-US", CaseClass.VIEWPORT_MOBILE, Outcome.COMPLETED, True),  # ran mobile
    ("mobile", "en-US", CaseClass.VIEWPORT_MOBILE, Outcome.NOT_RUN, True),  # never ran
    ("desktop", "hi-IN", CaseClass.HAPPY, Outcome.COMPLETED, True),  # locale never enacted
    ("iPhone 12", "en-US", CaseClass.VIEWPORT_MOBILE, Outcome.COMPLETED, False),
    ("desktop", "en-GB", CaseClass.HAPPY, Outcome.COMPLETED, False),
    ("kiosk", "en-US", CaseClass.HAPPY, Outcome.COMPLETED, True),  # unknown device
])
def test_persona_alignment_refuses_claims_the_run_did_not_enact(
    tmp_path: Path, device, locale, case_class, outcome, refused,
) -> None:
    store = _seed(tmp_path, persona=_persona(device, locale), case_class=case_class,
                  outcome=outcome)
    case, result = store.list_cases()[0], store.load_results(RUN)[0]
    assert (alignment_refusal(store.get_user_persona("p1"), case, result) is not None) is refused
    provider = _mock(UXJudgment(findings=[_draft()]))
    report = run_ux_advisory(store, RUN, provider)
    if refused:
        assert _status(report) is UXCaseStatus.REFUSED_CONDITION
        assert report.cases[0].findings == [] and provider.prompts == []
    else:
        assert _status(report) is UXCaseStatus.ANALYZED


# -- budget, opt-in, honesty ----------------------------------------------------------------
def test_persona_budget_skips_remaining_cases_and_counts_failed_calls(tmp_path: Path) -> None:
    store = _seed(tmp_path, max_calls=2, n_cases=4)
    provider = MockProvider(responses={"judge": [UXJudgment(findings=[_draft()])]})
    report = run_ux_advisory(store, RUN, provider)  # 2nd call: no queued answer -> ProviderError
    statuses = [c.status for c in report.cases]
    assert statuses == [UXCaseStatus.ANALYZED, UXCaseStatus.PROVIDER_ERROR,
                        UXCaseStatus.SKIPPED_BUDGET, UXCaseStatus.SKIPPED_BUDGET]
    assert report.calls_used == 2 and report.max_calls == 2


def test_persona_ux_is_opt_in_and_never_runs_when_disabled(tmp_path: Path) -> None:
    store = _seed(tmp_path, enabled=False)
    provider = _mock(UXJudgment())
    assert run_ux_advisory(store, RUN, provider) is None
    assert provider.prompts == [] and not store.paths.run_ux_report(RUN).exists()


def test_persona_dangling_or_missing_ref_is_recorded_not_inferred(tmp_path: Path) -> None:
    assert _status(run_ux_advisory(_seed(tmp_path / "a", ref="ghost"), RUN, _mock())) \
        is UXCaseStatus.MISSING_PERSONA
    assert _status(run_ux_advisory(_seed(tmp_path / "b", ref=None), RUN, _mock())) \
        is UXCaseStatus.NO_PERSONA


def test_persona_findings_citing_nonexistent_evidence_are_dropped(tmp_path: Path) -> None:
    store = _seed(tmp_path)
    answer = UXJudgment(findings=[_draft(step=1), _draft(step=9), _draft(step=None),
                                  _draft(step=None, path="ghost.png"),
                                  _draft(step=None, path="0.png", text="Icon has no label")])
    outcome = run_ux_advisory(store, RUN, _mock(answer)).cases[0]
    assert [f.finding for f in outcome.findings] == ["Label is unclear", "Icon has no label"]
    assert outcome.dropped_findings == 3


def test_persona_provider_failure_is_a_recorded_advisory_status(tmp_path: Path) -> None:
    class Boom(MockProvider):
        def judge(self, *a, **k):
            raise ProviderError("vendor said: secret detail")

    report = run_ux_advisory(_seed(tmp_path), RUN, Boom())
    assert _status(report) is UXCaseStatus.PROVIDER_ERROR
    assert "secret detail" not in (report.cases[0].reason or "")
